"""生成 16 套课程 PPT（统一视觉系统，可复现）。

设计约束（与天工看板同一套视觉语言）：
- 16:9，深蓝底 #0E1621，面板 #16222F，强调橙 #FF8A3D
- 每页只讲一件事；正文 ≤ 5 条，每条 ≤ 22 字（PPT 是提词器，不是讲义）
- 每页底部有课程/次课/页码，便于学生记笔记时定位

用法：
    python tools/build_decks.py                 # 全部 16 套
    python tools/build_decks.py --only S05      # 单套
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from deck_content import SESSIONS

COURSE = Path(__file__).resolve().parents[1]
OUT = COURSE / "slides"

BG = RGBColor(0x0E, 0x16, 0x21)
PANEL = RGBColor(0x16, 0x22, 0x2F)
PANEL2 = RGBColor(0x1B, 0x2A, 0x3A)
LINE = RGBColor(0x2A, 0x3D, 0x52)
TEXT = RGBColor(0xE8, 0xEE, 0xF4)
SUB = RGBColor(0x8F, 0xA5, 0xB8)
ACCENT = RGBColor(0xFF, 0x8A, 0x3D)
OK = RGBColor(0x2E, 0xD4, 0x7A)
WARN = RGBColor(0xFF, 0xB0, 0x20)
BAD = RGBColor(0xFF, 0x5C, 0x5C)
BLUE = RGBColor(0x3D, 0x9B, 0xFF)

FONT = "Microsoft YaHei"
W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.75)


def new_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    return prs


def _bg(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def _text(slide, left, top, width, height, text, size=18, color=TEXT,
          bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
          line_spacing=1.25):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.name = FONT
    return box


def _kicker(slide, text, color=ACCENT) -> None:
    _text(slide, MARGIN, Inches(0.55), Inches(6), Inches(0.4), text,
          size=13, color=color, bold=True)


def _title(slide, text, size=32, top=Inches(0.95), width=None) -> None:
    _text(slide, MARGIN, top, width or (W - 2 * MARGIN), Inches(1.2), text,
          size=size, color=TEXT, bold=True, line_spacing=1.15)


def _rule(slide, top=Inches(2.15), width=None, color=ACCENT) -> None:
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, top,
                                 Inches(1.1), Pt(4))
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()


def _footer(slide, s: dict, idx: int, total: int) -> None:
    _text(slide, MARGIN, H - Inches(0.55), Inches(8), Inches(0.3),
          f"智能体工程化项目实战 · {s['no']} {s['title']}",
          size=10, color=SUB)
    _text(slide, W - MARGIN - Inches(1.2), H - Inches(0.55), Inches(1.2),
          Inches(0.3), f"{idx} / {total}", size=10, color=SUB,
          align=PP_ALIGN.RIGHT)


def _bullets(slide, items, top=Inches(2.45), size=19, color=TEXT,
             marker="▸ ", gap=Inches(0.72), left=MARGIN, width=None):
    width = width or (W - 2 * MARGIN)
    for i, item in enumerate(items):
        y = top + gap * i
        _text(slide, left, y, Inches(0.35), Inches(0.5), marker,
              size=size, color=ACCENT, bold=True)
        _text(slide, left + Inches(0.42), y, width - Inches(0.42),
              Inches(0.6), item, size=size, color=color)


def _card(slide, left, top, width, height, title, items,
          accent=ACCENT, size=15):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top,
                                 width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = PANEL
    box.line.color.rgb = LINE
    box.line.width = Pt(1)
    box.adjustments[0] = 0.04
    box.shadow.inherit = False

    _text(slide, left + Inches(0.35), top + Inches(0.3),
          width - Inches(0.7), Inches(0.4), title, size=size + 3,
          color=accent, bold=True)
    for i, it in enumerate(items):
        _text(slide, left + Inches(0.35), top + Inches(0.85) + Inches(0.5) * i,
              width - Inches(0.7), Inches(0.45), f"· {it}", size=size,
              color=TEXT)


def _table(slide, headers, rows, top=Inches(2.4), size=14,
           col_widths=None):
    n_rows, n_cols = len(rows) + 1, len(headers)
    height = Inches(0.5) * n_rows
    left = MARGIN
    width = W - 2 * MARGIN
    shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    table = shape.table
    if col_widths:
        total = sum(col_widths)
        for i, cw in enumerate(col_widths):
            table.columns[i].width = Emu(int(width * cw / total))
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = ""
        p = cell.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = h
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = ACCENT
        r.font.name = FONT
        cell.fill.solid()
        cell.fill.fore_color.rgb = PANEL2
    for r_i, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.cell(r_i, c)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            run = p.add_run()
            run.text = str(val)
            run.font.size = Pt(size)
            run.font.color.rgb = TEXT
            run.font.name = FONT
            cell.fill.solid()
            cell.fill.fore_color.rgb = PANEL if r_i % 2 else BG
    return table


# ---------------------------------------------------------------- 页面模板
def cover(prs, s: dict) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0),
                                 Inches(0.18), H)
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()

    _text(slide, Inches(1.1), Inches(1.0), Inches(8), Inches(0.4),
          "智能体工程化项目实战 · 64 课时", size=15, color=SUB)
    _text(slide, Inches(1.1), Inches(1.55), Inches(10.5), Inches(0.5),
          f"{s['no']} · 第 {int(s['no'][1:])} 次课", size=16, color=ACCENT,
          bold=True)
    _text(slide, Inches(1.1), Inches(2.15), Inches(10.8), Inches(1.6),
          s["title"], size=44, color=TEXT, bold=True, line_spacing=1.1)
    _text(slide, Inches(1.1), Inches(3.9), Inches(10.8), Inches(0.6),
          s["sub"], size=20, color=SUB)
    _text(slide, Inches(1.1), Inches(5.4), Inches(10.8), Inches(0.4),
          f"课时 {s['hours']} · 布鲁姆层级 {s['bloom']} · 双元制（学校 + 企业）",
          size=13, color=SUB)
    _text(slide, Inches(1.1), Inches(5.9), Inches(10.8), Inches(0.4),
          "天工 Tiangong-EWA 工程基线 v1.2", size=12, color=LINE)


def quote(prs, s: dict, idx: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)
    _kicker(slide, "本课之问")
    _text(slide, MARGIN, Inches(2.4), W - 2 * MARGIN, Inches(2.2),
          s["hook"], size=34, color=TEXT, bold=True, line_spacing=1.3)
    _rule(slide, Inches(4.9))
    _text(slide, MARGIN, Inches(5.15), Inches(9), Inches(0.5),
          "先别急着找答案——这节课结束时，你要能自己讲清楚。",
          size=16, color=SUB)
    _footer(slide, s, idx, total)


def bullets_page(prs, s: dict, idx: int, total: int, kicker: str, title: str,
                 items, color=TEXT, marker="▸ ") -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)
    _kicker(slide, kicker)
    _title(slide, title, size=28)
    _rule(slide, Inches(1.95))
    _bullets(slide, items, top=Inches(2.4), color=color, marker=marker)
    _footer(slide, s, idx, total)


def situation_page(prs, s: dict, idx: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)
    _kicker(slide, "企业情境 · Lernsituation")
    _title(slide, "真实场景里，问题长这样", size=26)
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, MARGIN,
                                 Inches(2.2), W - 2 * MARGIN, Inches(2.7))
    box.fill.solid()
    box.fill.fore_color.rgb = PANEL
    box.line.color.rgb = ACCENT
    box.line.width = Pt(1.5)
    box.adjustments[0] = 0.05
    box.shadow.inherit = False
    for i, line in enumerate(s["situation"]):
        _text(slide, MARGIN + Inches(0.5), Inches(2.55) + Inches(0.62) * i,
              W - 2 * MARGIN - Inches(1.0), Inches(0.55), line, size=19,
              color=TEXT)
    _text(slide, MARGIN, Inches(5.3), Inches(10), Inches(0.5),
          "你的角色：准工程师。你的产出要被真实的人验收。", size=15,
          color=SUB)
    _footer(slide, s, idx, total)


def task_page(prs, s: dict, idx: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)
    _kicker(slide, "任务书 · Arbeitsauftrag")
    _title(slide, "这节课你要交付什么", size=28)
    _rule(slide, Inches(1.95))
    for i, t in enumerate(s["task"]):
        t = re.sub(r"^[ABC]\s+", "", t)
        y = Inches(2.35) + Inches(0.85) * i
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, MARGIN, y,
                                       Inches(0.5), Inches(0.5))
        badge.fill.solid()
        badge.fill.fore_color.rgb = ACCENT
        badge.line.fill.background()
        badge.shadow.inherit = False
        _text(slide, MARGIN, y + Inches(0.08), Inches(0.5), Inches(0.35),
              str(i + 1), size=15, color=BG, bold=True, align=PP_ALIGN.CENTER)
        _text(slide, MARGIN + Inches(0.75), y + Inches(0.05),
              W - 2 * MARGIN - Inches(0.75), Inches(0.6), t, size=19, color=TEXT)
    _text(slide, MARGIN, Inches(5.9), Inches(11), Inches(0.4),
          "纪律：无证据（截图/数据/代码）按“部分完成”计。", size=14, color=WARN)
    _footer(slide, s, idx, total)


def cards_page(prs, s: dict, idx: int, total: int, kicker: str, title: str,
               left: tuple, right: tuple) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)
    _kicker(slide, kicker)
    _title(slide, title, size=27)
    _rule(slide, Inches(1.95))
    half = (W - 2 * MARGIN - Inches(0.4)) / 2
    _card(slide, MARGIN, Inches(2.25), half, Inches(3.4), left[0], left[1],
          accent=ACCENT)
    _card(slide, MARGIN + half + Inches(0.4), Inches(2.25), half, Inches(3.4),
          right[0], right[1], accent=BLUE)
    _footer(slide, s, idx, total)


def steps_page(prs, s: dict, idx: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)
    _kicker(slide, "动手 · Ausführen")
    _title(slide, "按这四步做，别跳步", size=28)
    _rule(slide, Inches(1.95))
    n = len(s["steps"])
    gap = (W - 2 * MARGIN) / n
    for i, step in enumerate(s["steps"]):
        x = MARGIN + gap * i
        num = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, Inches(2.35),
                                     Inches(0.55), Inches(0.55))
        num.fill.solid()
        num.fill.fore_color.rgb = ACCENT
        num.line.fill.background()
        num.shadow.inherit = False
        _text(slide, x, Inches(2.42), Inches(0.55), Inches(0.4), str(i + 1),
              size=16, color=BG, bold=True, align=PP_ALIGN.CENTER)
        _text(slide, x, Inches(3.1), gap - Inches(0.3), Inches(1.4), step,
              size=15, color=TEXT)
        if i < n - 1:
            _text(slide, x + gap - Inches(0.35), Inches(2.42), Inches(0.4),
                  Inches(0.4), "→", size=20, color=LINE)
    _text(slide, MARGIN, Inches(5.2), Inches(11), Inches(0.4),
          "卡住了先看报错第一行；还不行再看 docs/07 环境问题速查。",
          size=14, color=SUB)
    _footer(slide, s, idx, total)


def end_page(prs, s: dict, idx: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)
    _kicker(slide, "小结与作业")
    _title(slide, "带走这三句话", size=28)
    _rule(slide, Inches(1.95))
    _bullets(slide, s["summary"], top=Inches(2.3), size=18, marker="✓ ",
             color=OK)
    _text(slide, MARGIN, Inches(4.75), Inches(11), Inches(0.4), "课后作业",
          size=16, color=ACCENT, bold=True)
    for i, hw in enumerate(s["homework"]):
        _text(slide, MARGIN, Inches(5.2) + Inches(0.42) * i, Inches(11),
              Inches(0.4), f"· {hw}", size=15, color=TEXT)
    _footer(slide, s, idx, total)


def build(s: dict) -> Path:
    prs = new_deck()
    total = 12
    cover(prs, s)
    quote(prs, s, 2, total)
    bullets_page(prs, s, 3, total, "学习目标 · Lernziele", "这节课结束时，你能",
                 s["objectives"])
    situation_page(prs, s, 4, total)
    task_page(prs, s, 5, total)
    bullets_page(prs, s, 6, total, "核心概念 ①", s["k1"][0], s["k1"][1])
    bullets_page(prs, s, 7, total, "核心概念 ②", s["k2"][0], s["k2"][1])
    bullets_page(prs, s, 8, total, "关键对比 · 想清楚再动手", s["compare"][0],
                 s["compare"][1], color=TEXT, marker="· ")
    bullets_page(prs, s, 9, total, "常见错误 · 别踩这些坑", "这些做法会被现场打回",
                 s["errors"], color=BAD, marker="✗ ")
    steps_page(prs, s, 10, total)
    bullets_page(prs, s, 11, total, "费曼检验 · 讲不清就是没学会",
                 "合上讲义，用三分钟讲给同伴", s["feynman"],
                 color=TEXT, marker="? ")
    end_page(prs, s, 12, total)

    name = f"{s['no']}_{s['title'].replace(' ', '').replace('/', '_')}.pptx"
    path = OUT / name
    OUT.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--only", default="", help="只生成某次课，如 S05")
    args = ap.parse_args()
    todo = [s for s in SESSIONS if not args.only or s["no"] == args.only]
    if not todo:
        print(f"未找到 {args.only}")
        return 2
    for s in todo:
        p = build(s)
        print(f"  {p.name}  ({len(p.read_bytes()) // 1024} KB)")
    print(f"完成 {len(todo)} 套")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
